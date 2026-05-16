from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Routey_数据库信息详解_20260516.docx")
CONTENT_WIDTH = 6.5

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(0, 0, 0)
MUTED = RGBColor(85, 85, 85)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
CAUTION_FILL = "FFF4CC"
RISK_FILL = "FCE8E6"
OK_FILL = "E6F4EA"
BORDER = "B7C3D0"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_para_format(paragraph, before=0, after=6, line=1.25, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_para(doc, text: str = "", size=11, bold=False, color: RGBColor = INK, before=0, after=6, line=1.25, align=None):
    p = doc.add_paragraph()
    set_para_format(p, before=before, after=after, line=line, align=align)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph()
    if level == 1:
        set_para_format(p, before=18, after=10, line=1.25)
        size, color = 16, BLUE
    elif level == 2:
        set_para_format(p, before=14, after=7, line=1.25)
        size, color = 13, BLUE
    else:
        set_para_format(p, before=10, after=5, line=1.25)
        size, color = 12, DARK_BLUE
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True, color=color)
    return p


def add_page_break(doc):
    doc.add_page_break()


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_in: float):
    width = int(width_in * 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_in)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_table_fixed(table, widths: Sequence[float]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    set_table_borders(table)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def clear_cell(cell):
    for p in cell.paragraphs:
        p._element.getparent().remove(p._element)


def cell_text(cell, text: str, size=8.6, bold=False, color: RGBColor = INK):
    clear_cell(cell)
    chunks = str(text).split("\n")
    for i, chunk in enumerate(chunks):
        p = cell.add_paragraph()
        set_para_format(p, before=0, after=2 if i < len(chunks) - 1 else 0, line=1.15)
        r = p.add_run(chunk)
        set_run_font(r, size=size, bold=bold, color=color)


def add_table(doc, headers: Sequence[str], rows: Iterable[Sequence[str]], widths: Sequence[float], caption: str | None = None):
    if caption:
        add_para(doc, caption, size=9, color=MUTED, before=4, after=4)
    rows = list(rows)
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    set_table_fixed(table, widths)
    repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, HEADER_FILL)
        cell_text(cell, header, size=8.8, bold=True, color=DARK_BLUE)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            fill = None
            low = str(value)
            if "必须" in low or "风险" in low or "不一致" in low:
                fill = RISK_FILL
            elif "建议" in low or "待接入" in low or "未持久化" in low:
                fill = CAUTION_FILL
            elif "已接入" in low or "当前可用" in low:
                fill = OK_FILL
            if fill:
                set_cell_shading(cell, fill)
            cell_text(cell, str(value), size=8.2)
    return table


def add_bullet(doc, items: Sequence[str]):
    for item in items:
        p = doc.add_paragraph(style=None)
        set_para_format(p, before=0, after=4, line=1.25)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        r1 = p.add_run("• ")
        set_run_font(r1, size=10.5, color=INK)
        r2 = p.add_run(item)
        set_run_font(r2, size=10.5, color=INK)


def add_footer_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para_format(footer, after=0, line=1.0)
    r = footer.add_run("Page ")
    set_run_font(r, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    footer._p.append(fld_char1)
    footer._p.append(instr)
    footer._p.append(fld_char2)


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    for style_name in ("Normal",):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_format(header, after=0, line=1.0)
    run = header.add_run("Routey 数据库说明 | compact_reference_guide")
    set_run_font(run, size=9, color=MUTED)
    add_footer_page_number(section)
    return doc


def build_document():
    doc = setup_doc()

    add_para(doc, "Routey 数据库信息详解表", size=24, bold=True, color=INK, before=4, after=4)
    add_para(
        doc,
        "基于 2026-05-16 重新扫描当前项目文件生成；重点区分“当前代码实际使用的存储”与“SQL 设计目标”。",
        size=12,
        color=MUTED,
        after=12,
    )
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["仓库路径", "/Users/zhangchenyuan/Desktop/Routey"],
            ["采用样式", "compact_reference_guide；Letter 纸张、1in 边距、固定宽表格"],
            ["扫描重点文件", "database-schema.sql、db/001_route_database.sql、src/server.ts、src/lib/tripRepository.ts、src/lib/tripStore.ts、src/lib/ai.ts、src/lib/amap-api.ts、routey-n8n-workflow.json、.env.example、deploy/aliyun/*、.routey-data/route-database.json"],
            ["一句话结论", "项目已经有“可运行的轻量路线库”，但数据库设计、n8n 工作流和后端 Repository 之间还没有完全统一；上线前必须合并成一份权威 migration。"],
        ],
        [1.6, 4.9],
    )

    add_heading(doc, "一、扫描范围与核心结论", 1)
    add_table(
        doc,
        ["文件", "角色", "当前读到的信息", "状态判断"],
        [
            ["database-schema.sql", "完整数据库设计", "包含 users、explore_routes、routes、itinerary_days、itinerary_places、user_trips、public_routes、route_segments、import_jobs、quiz_sessions、user_likes、spot_image_cache、destinations、route_profiles、workflow_logs、travel_info_cache、source_references，以及 2 个统计视图和 updated_at 触发器。", "设计最完整；但 routes 表缺少 n8n/后端正在使用的 city、route_theme、mood、summary、raw_json、source_url、likes 等字段，必须对齐。"],
            ["db/001_route_database.sql", "旧版/基础 migration", "包含 PostGIS、pg_trgm、routes、user_trips、route_days、route_spots、route_sources、route_media、route_segments；使用 GEOGRAPHY(POINT, 4326) 支持地理索引。", "和 database-schema.sql 存在 routes.id 类型与表拆分差异，不能无脑同时执行。"],
            ["src/lib/tripRepository.ts", "当前主数据访问层", "自动选择 PostgreSQL 或本地 JSON；PostgreSQL 只初始化 user_trips、routes、route_segments 三张表；保存完整 Trip JSON，公共路线也写入 routes。", "当前可用；但还没有落到完整设计里的 public_routes、import_jobs、spot_image_cache 等表。"],
            ["src/server.ts", "API 与业务编排", "提供 trips、public routes、imports、travel-info、explore、quiz、n8n-routes 等接口；另有 n8n PostgreSQL Pool 读取 routes/itinerary_days/itinerary_places。", "功能覆盖较全；但 import/quiz/image cache 仍是内存 TTL，重启丢失。"],
            ["routey-n8n-workflow.json", "自动生成精品路线", "每日 8:00 随机目的地，调用 DeepSeek，要求每天 5-8 个地点、总地点 >=12、覆盖景点/美食/购物/住宿/休闲，然后写 routes、itinerary_days、itinerary_places。", "路线质量规则写在 prompt 里，但来源真实性没有数据库级校验。"],
            [".routey-data/route-database.json", "本地 fallback 数据库", "当前包含 48 个 userTrips、3 个 publicRoutes、154 个 routeSegments。", "开发预览可用；上线不能依赖本地 JSON。"],
            [".env.example", "环境变量模板", "已有 VITE_TENCENT_MAP_KEY、DEEPSEEK_API_KEY、AMAP_WEB_KEY、ENABLE_LIVE_ROUTE_PLANNING。", "缺少 DATABASE_URL、POSTGRES_URL、N8N_DATABASE_URL、PGSSLMODE、OSS 配置示例，建议补齐。"],
            ["deploy/aliyun/*", "阿里云部署脚本", "Nginx 反代 127.0.0.1:3001，PM2 启动 scripts/serve-aliyun.mjs。", "Node 部署链路清楚；还缺 RDS 初始化、环境变量注入、OSS 配置说明。"],
        ],
        [1.35, 1.15, 2.65, 1.35],
    )

    add_heading(doc, "二、当前数据架构总览", 1)
    add_table(
        doc,
        ["层级", "当前实现", "数据库落点", "说明"],
        [
            ["前端状态", "React Query hooks：useTrips、useTrip、useExploreRoutes、usePublicRoutes、useTravelInfo、useFeaturedRoutes。", "不直连数据库，只访问 /api/*。", "这是正确方向，未来切 DB 不影响大部分 UI。"],
            ["Repository 主线", "getTripRepository(env) 读取 DATABASE_URL/POSTGRES_URL；存在则走 PostgreSQL，否则走 FileTripRepository。", "PostgreSQL: user_trips、routes、route_segments；本地: .routey-data/route-database.json。", "当前核心行程 CRUD 已统一在 Repository。"],
            ["用户私有行程", "导入、保存公共路线、Explore 添加后生成个人 Trip。", "user_trips.trip_json；本地 JSON 的 userTrips。", "默认 private_draft。"],
            ["公共路线库", "publishTrip 会复制一份 public route。", "当前写入 routes 表，而不是 database-schema.sql 里的 public_routes 表。", "设计表和实现表需要统一。"],
            ["n8n 精品路线", "独立 n8n pool 查询 status='published' 的路线。", "n8n 期望 routes、itinerary_days、itinerary_places。", "这条线和 Repository 的 routes 表同名但结构不同，是最大冲突点。"],
            ["路线交通缓存", "前端请求 /api/travel-info；先查内存 TTL，再查 Repository，再生成。", "route_segments.data；内存 travelCache。", "ENABLE_LIVE_ROUTE_PLANNING=false 时生成估算路线并缓存。"],
            ["图片缓存", "服务端 /api/spot-image 先 Bing CN，再 Wikipedia，再 fallback。", "当前仅 spotImageCache 内存；SQL 已设计 spot_image_cache。", "上线建议持久化，减少外部请求和不稳定。"],
            ["AI 解析/Quiz", "DeepSeek 生成结构化 Trip，坐标缺失时 Nominatim 补点。", "结果写 user_trips；任务状态仍在 TTLCache。", "import_jobs、quiz_sessions 暂未接入。"],
        ],
        [1.25, 2.0, 1.75, 1.5],
    )

    add_heading(doc, "三、数据库表清单", 1)
    add_table(
        doc,
        ["对象", "来源 SQL", "用途", "关键关系/索引", "当前代码使用情况"],
        [
            ["users", "database-schema.sql", "用户主体：匿名、手机号、微信、邮箱，保存 preferences。", "phone/wechat 唯一索引；被 owner_id 引用。", "设计表；当前代码只有 anonymous_user_id 字符串，未真正创建/查询 users。"],
            ["explore_routes", "database-schema.sql", "精选路线元数据：目的地、天数、来源、质量分、封面、标签。", "destination/region/days/tags/is_active 索引。", "设计表；当前 Explore 仍主要来自 tripPlanner/extraRoutes 硬编码。"],
            ["routes", "database-schema.sql + Repository + n8n", "名称冲突核心表：设计中承载 n8n/公共路线；Repository 中承载 public Trip JSON；n8n 中承载 route_title/city/raw_json。", "idx_routes_*；n8n days route_id 外键。", "当前必须重构统一，否则同名 routes 表会互相打架。"],
            ["itinerary_days", "database-schema.sql + n8n", "n8n 路线的 Day 级数据。", "route_id -> routes.id；idx_days_route。", "server.ts /api/explore、/api/explore/add、/api/n8n-route-detail 已读取。"],
            ["itinerary_places", "database-schema.sql + n8n", "n8n Day 下的地点、餐厅、酒店、购物、休闲等 POI。", "day_id -> itinerary_days.id；category/geo 索引。", "server.ts 已读取并转成 Trip.days[].spots。"],
            ["user_trips", "database-schema.sql + Repository", "用户私人行程副本，保存完整 trip_json。", "owner/visibility、favorite、destination、tags 索引。", "当前主表；GET /api/trips、导入、保存路线、编辑都依赖它。"],
            ["public_routes", "database-schema.sql", "用户发布/官方路线的公共库。", "visibility+quality、destination/country、tags、day_count 索引。", "设计表；当前 Repository 还未使用，公共路线写在 routes。"],
            ["route_segments", "database-schema.sql + Repository", "缓存两个景点之间的距离、时间、mode、polyline。", "trip_id/day_id 索引。", "当前已接入 /api/travel-info；本地 JSON 有 154 条。"],
            ["import_jobs", "database-schema.sql", "AI 导入任务状态、进度、生成 trip_id、错误信息。", "owner/status 索引。", "设计表；当前 importJobs 是内存 TTL，重启丢失。"],
            ["quiz_sessions", "database-schema.sql", "Quiz 问卷、匹配路线、AI 生成结果。", "owner 索引。", "设计表；当前 quizAiJobs 是内存 TTL，未持久化。"],
            ["user_likes", "database-schema.sql", "用户对 explore_route/public_route 收藏或点赞。", "owner、target 索引；owner+target 唯一。", "设计表；当前 favorite 在 Trip JSON 内，公共点赞未入库。"],
            ["spot_image_cache", "database-schema.sql", "景点图片查询缓存。", "expires_at 索引。", "设计表；当前 spotImageCache 内存缓存。"],
            ["destinations", "database-schema.sql", "目的地基础信息、签证、货币、语言、封面、排序。", "region/sort_order 索引。", "设计表；当前 /api/destinations 从 tripPlanner 获取。"],
            ["route_profiles", "database-schema.sql", "一键添加时用的路线模板 Day[]。", "destination+profile_key 唯一。", "设计表；当前模板在 tripPlanner.ts。"],
            ["workflow_logs", "database-schema.sql", "记录 n8n 工作流每一步输入输出、耗时、错误。", "route/status 索引。", "设计表；n8n workflow 还未写日志表。"],
            ["travel_info_cache", "database-schema.sql", "高德/OSRM 交通缓存，以坐标对为 key。", "expires_at 索引。", "设计表；当前缓存实际在 route_segments 和内存。"],
            ["source_references", "database-schema.sql", "精选路线来源验证记录。", "explore_route_id -> explore_routes.id。", "设计表；真实性验证下一步应落这里。"],
            ["v_destination_stats", "database-schema.sql", "目的地路线数、平均质量分、天数范围、总点赞。", "视图基于 explore_routes。", "设计视图；需迁移 explore_routes 后才有价值。"],
            ["v_user_trip_stats", "database-schema.sql", "用户总行程、进行中、完成、草稿、收藏统计。", "视图基于 user_trips。", "上线后可用于 Profile/运营看板。"],
            ["route_days / route_spots / route_sources / route_media", "db/001_route_database.sql", "旧版更规范的拆分结构，route_spots 使用 PostGIS 点位，route_media 对接 OSS。", "route_spots.location GIST；route_media 保存 oss_key/url。", "当前代码未接入；但可作为 v2 规范化方向。"],
        ],
        [1.25, 1.05, 1.65, 1.4, 1.15],
    )

    add_heading(doc, "四、字段级说明：核心业务表", 1)
    add_table(
        doc,
        ["表", "关键字段", "字段解释", "写入来源", "读取/消费位置"],
        [
            ["user_trips", "id, owner_id, name, date, cover, status, favorite, visibility, source_route_id, destination, country, tags, quality_score, trip_json, created_at, updated_at", "用户私人行程的唯一可信副本。trip_json 保存完整前端 Trip 对象，方便快速迭代；destination/country/tags/quality_score 是检索和排序冗余字段。", "repository.saveUserTrip；AI 导入；Explore 添加；公共路线保存；Quiz AI。", "GET /api/trips、GET/PATCH/DELETE /api/trips/:id、地图页、我的行程。"],
            ["routes（Repository 视角）", "id TEXT, owner_id, source_trip_id, visibility, destination, country, tags, day_count, quality_score, trip_json, published_at", "当前公共路线库实际表。每条公共路线也是完整 Trip JSON；visibility='public' 才返回。", "publishTrip；初次空库 seed。", "GET /api/routes/public、POST /api/routes/:id/save、GET /api/trips/:id fallback。"],
            ["route_segments", "id, trip_id, day_id, segment_key, data, updated_at", "缓存某个 Trip 某天所有相邻 POI 的交通段。data 为 key=>TravelInfo，包含 distance、duration、mode、polyline。", "/api/travel-info 在生成或估算后保存。", "地图页连线、距离/时间展示；先查内存，再查 DB。"],
            ["routes（n8n 视角）", "id SERIAL, route_title, city, country, destination, route_theme, mood, summary, raw_json, source_url, status, quality_score, best_time, days_count, budget_level, travel_type, pace, tags, cover_url, likes", "n8n 自动生成路线的主表，status='published' 才展示。注意：这些字段没有全部出现在 database-schema.sql 当前 routes 定义里。", "routey-n8n-workflow.json 的 Postgres Routes 节点。", "/api/explore、/api/n8n-routes、/api/n8n-route-detail、/api/explore/add。"],
            ["itinerary_days", "id, route_id, day_number, title, summary, route_desc, estimated_budget", "n8n 路线按天拆分。route_desc 对应前端 Day.route；summary 用于列表和详情。", "n8n Postgres1 Days 节点。", "n8nRouteToTrip 转成 Trip.days。"],
            ["itinerary_places", "id, day_id, place_name, visit_time, duration_min, description, intro, tags, place_order, latitude, longitude, category, rating, price, image_url, address, phone, opening_hours, external_id", "地点级信息，覆盖景点/美食/购物/住宿/休闲。坐标用于地图点和路线段生成；category/rating/price 用于详情弹窗和筛选。", "n8n Postgres2 Places 节点。", "n8nRouteToTrip 转成 Spot；地图页、路线详情页、详情弹窗。"],
            ["explore_routes", "id, title, destination, region, days, spots, source, source_name, source_url, source_verified, source_references, quality_score, likes, cover, tags, includes, profile_key", "精选路线卡片和来源真实性的元数据表。source_references JSONB 可临时保存引用，但长期建议拆到 source_references。", "未来 seed/migration 或运营后台。", "未来替代 tripPlanner.ts/extraRoutes.ts；/api/explore 应从这里读。"],
            ["public_routes", "id, owner_id, source_trip_id, destination, country, tags, day_count, quality_score, visibility, trip_json, published_at", "设计上的公共路线库。和当前 Repository 的 routes 表职责重复，建议二选一或改名。", "未来审核通过后写入。", "未来 GET /api/routes/public、保存公共路线。"],
            ["source_references", "id, explore_route_id, title, publisher, url, verified_at", "用于支撑“路线来源真实性”。应保存官方旅游局、景点官网、酒店/餐厅官网、可信攻略平台链接。", "未来爬虫、运营录入、AI 生成后的验证流程。", "Explore 详情、路线来源展示、审核后台。"],
            ["route_media", "id, route_id, user_trip_id, spot_id, media_type, oss_key, url, width, height", "OSS 图片/截图/原始导入内容索引表。数据库只存 key 和可访问 URL，不存二进制。", "未来 OSS 上传服务。", "封面图、地点图、导入截图回溯。"],
        ],
        [1.1, 1.65, 1.8, 1.0, 0.95],
    )

    add_heading(doc, "五、API 到数据库映射", 1)
    add_table(
        doc,
        ["接口", "功能", "当前数据表/缓存", "持久化情况", "备注"],
        [
            ["GET /api/trips", "返回当前用户私有行程", "user_trips 或本地 JSON userTrips", "已接入持久化", "ownerId 来自 x-routey-user-id，默认 anonymous_user_id。"],
            ["GET /api/trips/:id", "读取单个行程", "先 user_trips，再 public routes/routes", "已接入持久化", "公共路线也可被详情页读取。"],
            ["PATCH /api/trips/:id", "收藏、重排、优化、增删改 spot", "user_trips.trip_json", "已接入持久化", "public 路线不可直接编辑。"],
            ["DELETE /api/trips/:id", "删除个人行程", "user_trips", "已接入持久化", "物理删除；后续可改 archived。"],
            ["POST /api/imports", "AI 解析链接/文本/图片/视频", "importJobs 内存 + user_trips", "行程已持久化，任务状态未持久化", "建议接入 import_jobs。"],
            ["GET /api/imports/:id", "轮询解析进度", "TTLCache importJobs", "未持久化", "服务重启后任务状态丢失。"],
            ["GET /api/routes/public", "公共路线库列表", "Repository routes 表或 JSON publicRoutes", "已接入持久化", "没有使用 public_routes 设计表。"],
            ["POST /api/trips/:id/publish", "提交/发布个人行程", "user_trips + routes", "已接入持久化", "当前会马上生成 public route；审核状态还不完整。"],
            ["POST /api/routes/:id/save", "保存公共路线到个人", "routes -> user_trips", "已接入持久化", "复制新私人副本，不影响原公共路线。"],
            ["POST /api/travel-info", "生成/读取每天路线段", "travelCache + route_segments", "已接入持久化", "默认估算路线；ENABLE_LIVE_ROUTE_PLANNING=true 才调高德/OSRM。"],
            ["GET /api/spot-image", "按地点找高清图", "spotImageCache 内存", "未持久化", "SQL 已有 spot_image_cache，建议接入。"],
            ["GET /api/destinations", "目的地分组", "tripPlanner 硬编码", "未接入 DB", "SQL 已有 destinations。"],
            ["GET /api/explore", "热门路线列表", "硬编码 getExploreRoutes + n8n routes/days/places 聚合", "部分 DB", "n8n 线路会排在硬编码前。"],
            ["POST /api/explore/add", "添加热门路线到我的行程", "n8n routes/days/places 或 tripPlanner/profile 或 AI job", "最终写 user_trips", "如果 routeId 是 n8n-*，直接从 DB 组装 Trip。"],
            ["POST /api/quiz-recommend", "问卷匹配和 AI 路线生成", "quizAiJobs 内存 + user_trips", "结果持久化，任务状态未持久化", "建议接入 quiz_sessions。"],
            ["GET /api/n8n-routes", "首页/推荐展示 n8n 路线", "n8n routes", "依赖 n8n DB", "按 quality_score、created_at 排序。"],
            ["GET /api/n8n-route-detail", "n8n 路线详情", "routes + itinerary_days + itinerary_places", "依赖 n8n DB", "返回 raw route 和转换后的 Trip。"],
        ],
        [1.35, 1.35, 1.35, 1.15, 1.3],
    )

    add_heading(doc, "六、关键数据流", 1)
    add_table(
        doc,
        ["流程", "输入", "处理步骤", "最终入库", "当前风险/建议"],
        [
            ["用户导入攻略", "链接、文本、截图、视频描述", "POST /api/imports -> DeepSeek parseWithAI -> 坐标补齐 -> optimizeTrip -> saveUserTrip。", "user_trips.trip_json；任务状态在内存。", "建议把 import_jobs 落库，并保存原始内容到 OSS/route_sources。"],
            ["n8n 自动精品路线", "目的地与关键词", "Schedule -> DeepSeek -> Parse Route -> Postgres routes -> split days -> itinerary_days -> split places -> itinerary_places。", "routes、itinerary_days、itinerary_places。", "必须补齐 routes 表字段；source_url 目前多为空，真实性需要 source_references。"],
            ["Explore 添加路线", "routeId", "n8n-* 直接 DB 组装；硬编码 profile 直接生成；无 profile 时触发 AI job。", "user_trips。", "建议迁移硬编码 explore/routes/profile 到 DB，统一来源。"],
            ["公共路线发布", "个人 user_trip id", "publishTrip 复制 public route，同时把个人行程 visibility 改为 published_pending。", "user_trips + routes。", "当前缺审核工作流，public_routes 设计表未使用。"],
            ["地图路线段", "tripId/dayId", "读取 day.spots -> 查内存 -> 查 route_segments -> getDayTravelInfo -> 保存。", "route_segments.data。", "默认不调高德，只生成估算 polyline；这符合低成本方案。"],
            ["景点图片", "地点查询词", "Bing CN 优先，Wikipedia 备用，最后 fallback 图片。", "当前仅内存。", "建议落 spot_image_cache，且把长期图片转存 OSS，避免外链失效。"],
        ],
        [1.15, 1.1, 2.05, 1.05, 1.15],
    )

    add_heading(doc, "七、环境变量与云资源", 1)
    add_table(
        doc,
        ["配置", "当前代码读取位置", "用途", "上线建议"],
        [
            ["DATABASE_URL", "tripRepository.ts、server.ts n8n fallback", "主 PostgreSQL 连接；Repository 有则走 PostgreSQL。", "必须加入 .env.example 和阿里云 PM2 环境。"],
            ["POSTGRES_URL", "tripRepository.ts", "DATABASE_URL 的备用名称。", "二选一即可，建议统一 DATABASE_URL。"],
            ["N8N_DATABASE_URL", "server.ts getN8nPool", "n8n 生成路线库连接。", "如果与主库共用，显式配置同一个 RDS；如果拆库，保证 schema 一致。"],
            ["PGSSLMODE", "tripRepository.ts", "控制 pg Pool SSL；PGSSLMODE=disable 时不启用 SSL。", "阿里云 RDS 公网/生产建议启用 SSL 或走内网白名单。"],
            ["DEEPSEEK_API_KEY", "ai.ts、n8n 凭证", "AI 解析、Quiz、路线生成。", "只放服务端或 n8n credential，不能进前端 bundle。"],
            ["VITE_TENCENT_MAP_KEY", ".env.example、前端地图", "腾讯地图 JS API GL 展示 key。", "前端公开 key，需配置域名白名单。"],
            ["AMAP_WEB_KEY", "amap-api.ts", "国内实时路线规划/地理服务。", "默认可不填；ENABLE_LIVE_ROUTE_PLANNING=false 时不会调用路线规划。"],
            ["ENABLE_LIVE_ROUTE_PLANNING", "amap-api.ts", "是否启用高德/OSRM 实时路线。", "低成本方案保持 false；生成估算路线并缓存。"],
            ["OSS_*（未出现）", "当前未接入", "未来保存导入原始文件、封面、地点图。", "建议增加 OSS_REGION、OSS_BUCKET、OSS_ACCESS_KEY_ID、OSS_ACCESS_KEY_SECRET、OSS_PUBLIC_BASE_URL。"],
        ],
        [1.35, 1.35, 1.45, 2.35],
    )

    add_heading(doc, "八、当前必须处理的数据库问题", 1)
    add_table(
        doc,
        ["优先级", "问题", "影响", "建议处理"],
        [
            ["P0", "routes 表存在三种含义：完整 SQL 的 n8n/公共混合表、Repository 的 public Trip JSON 表、n8n workflow 的自动生成路线表。", "上线迁移会冲突，n8n insert 可能缺列失败，公共路线查询可能读错表。", "拆分命名：public_routes 存公共 Trip JSON；generated_routes 或 ai_routes 存 n8n；Repository 不再用 routes 作为公共表。"],
            ["P0", "database-schema.sql 的 routes 缺少 n8n workflow 正在插入的 city、route_theme、mood、places、summary、raw_json、source_url、likes。", "执行当前 schema 后，n8n Postgres Routes 节点会失败。", "补齐字段或改 n8n workflow 插入字段；两边只能保留一份权威定义。"],
            ["P0", ".env.example 缺少数据库相关变量。", "新环境无法知道要配置 DATABASE_URL/N8N_DATABASE_URL。", "补齐 DATABASE_URL、N8N_DATABASE_URL、PGSSLMODE、OSS_*。"],
            ["P1", "import_jobs、quiz_sessions、spot_image_cache、destinations、route_profiles、source_references 设计了但代码未接入。", "重启丢任务；首页/Explore 仍依赖硬编码；来源真实性无法沉淀。", "按“任务、图片缓存、目的地/路线种子、来源验证”的顺序接入。"],
            ["P1", "sourceVerified 在 n8n route 映射里直接设为 true，但没有真实 source_references。", "会让“所有路线来源真实”变成 UI 声明，而非可审计数据。", "只有存在至少 1 条 verified source reference 时才置 true。"],
            ["P1", "updated_at 触发器 DO 块重复执行可能因同名 trigger 已存在而失败。", "多次迁移不幂等。", "创建触发器前检查 pg_trigger 或先 DROP TRIGGER IF EXISTS。"],
            ["P2", "当前 Trip JSON 适合快速迭代，但难做复杂检索和去重。", "路线库变大后，目的地/景点/餐厅搜索会慢，质量治理难。", "v1 保留 trip_json，同时把 route_days/route_spots 或 itinerary_* 规范化为可检索表。"],
            ["P2", "本地 JSON 当前有 48 个用户行程，其中示例首条只有 9 个 spots，不符合精品路线标准。", "数据质量不统一。", "上线 seed/import 规则加入 min_spots、category coverage、坐标完整率、来源完整率校验。"],
        ],
        [0.65, 2.1, 1.65, 2.1],
    )

    add_heading(doc, "九、推荐的上线数据库形态", 1)
    add_bullet(
        doc,
        [
            "保留 PostgreSQL 作为唯一结构化数据库，阿里云 RDS PostgreSQL 承载用户、路线、POI、缓存、任务状态。",
            "保留 user_trips.trip_json 作为前端快速渲染副本，同时把公共/精品路线拆成 public_routes、route_days、route_spots 或 generated_routes、itinerary_days、itinerary_places。",
            "把硬编码 Explore 路线迁移进 explore_routes 和 route_profiles，前端继续走 /api/explore，不再知道数据来自代码还是 DB。",
            "OSS 只存文件和图片，数据库存 oss_key、url、width、height、source、expires_at。",
            "真实性用 source_references 做审核证据：每条公共精品路线至少有官方/可信来源，AI 生成内容只能作为草稿，不能自动标记 sourceVerified=true。",
            "路线质量分 quality_score 应由规则计算：地点数量、坐标完整率、类别覆盖、来源数量、图片覆盖、用户反馈。",
        ],
    )
    add_table(
        doc,
        ["模块", "v1 建议表", "说明"],
        [
            ["用户与权限", "users", "匿名 ID 先落 users，未来绑定手机号/微信/邮箱。"],
            ["私人行程", "user_trips", "用户可编辑副本，默认 private_draft。"],
            ["公共路线", "public_routes + source_references", "审核通过后 public；真实来源必须可追溯。"],
            ["AI/n8n 生成库", "generated_routes + generated_days + generated_places，或沿用 routes/itinerary_days/itinerary_places 但改名避免冲突", "自动路线先进入待审核池。"],
            ["Explore 官方库", "explore_routes + route_profiles + destinations", "替代硬编码目的地和路线卡片。"],
            ["地图交通", "route_segments 或 travel_info_cache", "按 trip/day 或坐标 pair 缓存估算/真实路线。"],
            ["媒体文件", "route_media + OSS", "封面、地点图、用户上传图片、原始攻略文件。"],
            ["异步任务", "import_jobs + quiz_sessions + workflow_logs", "导入、Quiz、n8n 执行都可恢复和排查。"],
        ],
        [1.25, 2.2, 3.05],
    )

    add_heading(doc, "十、下一步落地顺序", 1)
    add_table(
        doc,
        ["顺序", "任务", "验收标准"],
        [
            ["1", "合并 schema，产出唯一 migration：解决 routes 命名冲突和缺列问题。", "空 RDS 一键执行成功；n8n insert 成功；Repository CRUD 成功。"],
            ["2", "补齐环境变量模板和阿里云部署说明。", "新机器只按文档配置 env 即可启动服务并连 RDS。"],
            ["3", "把 import_jobs、quiz_sessions、spot_image_cache 从内存切到 DB。", "服务重启后导入进度、Quiz 状态、图片缓存不丢。"],
            ["4", "迁移 Explore 硬编码数据到 explore_routes/destinations/route_profiles。", "首页、热门目的地、国家路线页刷新后全部从 DB 渲染。"],
            ["5", "接入 source_references 和质量校验。", "每条公共精品路线都能展示来源；不满足 12+ 地点和类别覆盖的不允许 public。"],
            ["6", "接入 OSS route_media。", "导入截图、封面、地点详情图不依赖第三方外链。"],
        ],
        [0.65, 3.0, 2.85],
    )

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build_document()
