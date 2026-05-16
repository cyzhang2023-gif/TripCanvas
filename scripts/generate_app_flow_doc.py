from __future__ import annotations

from pathlib import Path

from docx.oxml.ns import qn

from generate_database_doc import (
    INK,
    MUTED,
    add_bullet,
    add_heading,
    add_para,
    add_table,
    set_para_format,
    set_run_font,
    setup_doc,
)


OUT = Path("Routey_APP全链路运行工作详解_20260516.docx")


def reset_header(doc):
    header = doc.sections[0].header.paragraphs[0]
    header._p.clear_content()
    set_para_format(header, after=0, line=1.0)
    run = header.add_run("Routey APP 全链路说明 | compact_reference_guide")
    set_run_font(run, size=9, color=MUTED)


def build_document():
    doc = setup_doc()
    reset_header(doc)

    add_para(doc, "Routey APP 全链路运行工作详解表", size=24, bold=True, color=INK, before=4, after=4)
    add_para(
        doc,
        "基于 2026-05-16 重新扫描当前前端、后端、AI、地图、数据层和部署入口生成；用于理解整个 APP 从启动、用户操作、接口处理到数据存储的完整链路。",
        size=12,
        color=MUTED,
        after=12,
    )
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["仓库路径", "/Users/zhangchenyuan/Desktop/Routey"],
            ["技术栈", "TanStack Start + React 19 + React Query + TanStack Router + Tailwind CSS + MapLibre GL + Node fetch handler + PostgreSQL/本地 JSON fallback"],
            ["运行入口", "开发：npm run dev；阿里云构建：npm run build:aliyun；生产启动：node scripts/serve-aliyun.mjs 或 npm run start:aliyun；PM2 使用 ecosystem.config.cjs。"],
            ["核心链路", "用户在移动端 UI 操作 -> React Query/hook 发起 /api 请求 -> src/server.ts 分发 -> Repository/AI/地图/图片/n8n 数据源处理 -> 返回 JSON -> 前端缓存刷新并跳转/重绘。"],
            ["一句话结论", "APP 已形成完整闭环：发现目的地、导入攻略、AI 解析、路线详情地图、编辑行程、我的行程、分享、外部导航都已串联；上线前重点是数据库统一、异步任务持久化和来源真实性审核。"],
        ],
        [1.45, 5.05],
    )

    add_heading(doc, "一、全链路总览", 1)
    add_table(
        doc,
        ["阶段", "入口文件/模块", "做什么", "输入", "输出/下一步"],
        [
            ["1. 访问 APP", "deploy/aliyun/nginx.conf -> scripts/serve-aliyun.mjs", "Nginx 把请求反代到 127.0.0.1:3001；Node 判断是否静态资源，否则转成 Fetch Request 交给 TanStack Start server entry。", "浏览器请求 URL、Header、Body。", "静态资源或 SSR/接口响应。"],
            ["2. 应用壳初始化", "src/routes/__root.tsx、src/router.tsx、src/start.ts", "创建 QueryClient、Router、RootShell、全局 meta/css、错误页；server middleware 捕获未处理异常并返回品牌错误页。", "TanStack Start 请求上下文。", "页面 Outlet + React Query Provider。"],
            ["3. 首页发现", "src/routes/index.tsx", "加载 n8n 推荐路线、今日灵感、推荐路线、用户正在看；支持搜索/粘贴链接/打开导入弹层/进入 Quiz。", "用户输入内容、n8n featured routes。", "跳 /parsing、/quiz、/destinations 或 /explore。"],
            ["4. 导入解析", "src/routes/import.tsx、src/routes/parsing.tsx、server.ts /api/imports", "用户提交链接/图片/文本/视频；服务端创建 import job，调用 DeepSeek 解析并保存 Trip；解析页轮询进度。", "SourceKind + content；可包含压缩后的 base64 截图。", "job done 后自动跳 /trip?id=...&focus=map。"],
            ["5. 目的地与热门路线", "src/routes/destinations.tsx、src/routes/explore.tsx、server.ts /api/destinations /api/explore", "目的地页按地区展示；热门路线页按分类、城市、天数、方式、排序筛选；添加路线到我的行程。", "dest 参数、筛选状态、routeId。", "n8n 路线直接组装 Trip；硬编码路线走 profile 或 AI job。"],
            ["6. Quiz 推荐", "src/routes/quiz.tsx、server.ts /api/quiz-recommend", "6 步问卷收集范围、风格、天数、同行人、预算、季节；返回 DB 匹配和 AI 生成任务。", "QuizAnswers。", "显示匹配卡片；AI 完成后可跳行程详情。"],
            ["7. 行程详情地图", "src/routes/trip.tsx、src/components/AMapView.tsx", "读取 Trip，展示 42vh 地图、Day chips、路线点线、地点卡、预算、编辑、详情弹窗、外部高德导航。", "tripId、dayId、spots、travel info。", "地图重绘、Trip PATCH、打开高德 URI、分享页。"],
            ["8. 我的行程", "src/routes/my-trips.tsx", "读取当前用户行程，支持全部/收藏/草稿过滤，收藏、重新生成、删除。", "GET /api/trips 返回 Trip[]。", "打开 /trip；PATCH/DELETE 后刷新列表。"],
            ["9. 分享", "src/routes/share.tsx", "生成分享卡和文本摘要，支持 Web Share、复制、下载 txt。", "tripId。", "本地分享/下载，不需要额外后端。"],
            ["10. 数据持久化", "src/lib/tripRepository.ts", "自动选择 PostgreSQL 或本地 .routey-data/route-database.json；统一保存用户行程、公共路线和路线段缓存。", "Trip、filters、TravelInfoMap。", "user_trips/routes/route_segments 或 JSON 文件。"],
        ],
        [0.75, 1.45, 2.0, 1.1, 1.2],
    )

    add_heading(doc, "二、页面与用户流程详表", 1)
    add_table(
        doc,
        ["页面", "用户看到/可做", "前端关键状态", "调用接口/动作", "成功后跳转或刷新"],
        [
            ["/", "首页：Hero 搜索/AI生成、四个功能入口、今日灵感、推荐路线、用户正在看、导入底部弹层。", "content、showImport、activeKind、drawerText、imageBase64、submitting、watchOffset。", "useFeaturedRoutes -> GET /api/n8n-routes?limit=30；createImport -> POST /api/imports。", "有输入则进入 /parsing；无输入点 AI 生成进入 /quiz；灵感/推荐进入 /explore。"],
            ["/import", "独立导入页：链接、截图、文本、视频四种入口。", "kind、content、fileName、error、submitting。", "useTripActions.createImport -> POST /api/imports。", "拿到 jobId 后跳 /parsing?jobId=...。"],
            ["/parsing", "AI 解析进度页：进度条、四步状态、错误重试。", "jobId、job、steps、progress。", "useImportJob 每 1500ms GET /api/imports/:id。", "job.status=done 且有 tripId 时 600ms 后 replace 到 /trip?id=tripId&focus=map。"],
            ["/destinations", "按地区展示目的地卡片，支持搜索国家/城市/灵感。", "search、filteredGroups。", "useDestinations -> GET /api/destinations。", "点击目的地跳 /explore?dest=...。"],
            ["/explore", "某国家热门路线页：Hero、分类 tab、筛选 pill、路线卡、本地人推荐。", "cat、city、days、mode、budget、sort、addingId、toast。", "useExploreRoutes -> GET /api/explore?dest=；addExploreRoute -> POST /api/explore/add。", "直接生成 Trip 则 /trip?id=...；需要 AI 生成则 /parsing?jobId=...。"],
            ["/quiz", "六步旅行偏好问卷与推荐结果。", "step、answers、loading、dbMatches、aiRoutes、addingRoute。", "POST /api/quiz-recommend；轮询 POST /api/quiz-ai-status；添加路线 POST /api/explore/add。", "匹配路线可加入行程；AI 完成后进入 /trip。"],
            ["/my-trips", "我的行程列表：全部、收藏、草稿；卡片上可收藏、重新生成、删除。", "trips、tab、filtered。", "GET /api/trips；PATCH /api/trips/:id；DELETE /api/trips/:id。", "编辑动作后 React Query invalidate；点击卡片进入 /trip。"],
            ["/trip", "行程详情：地图、路线天数、地点列表、预算、编辑、地点详情、外部导航。", "trip、activeDay、selectedSpot、detailSpot、showBudget、editing。", "GET /api/trips/:id；POST /api/travel-info；PATCH /api/trips/:id；GET /api/spot-image。", "地图/列表联动；编辑保存后刷新 Trip；分享按钮进入 /share。"],
            ["/share", "分享行程卡片，选择模板、复制/分享/下载摘要。", "template、summary。", "GET /api/trips/:id；浏览器 navigator.share/clipboard；Blob 下载。", "生成本地分享文本，不改数据库。"],
        ],
        [0.85, 1.65, 1.2, 1.5, 1.3],
    )

    add_heading(doc, "三、前端数据层与缓存", 1)
    add_table(
        doc,
        ["Hook/方法", "React Query key", "接口/动作", "缓存策略", "消费者"],
        [
            ["useTripsQuery", "['trips']", "GET /api/trips", "staleTime 10s；写操作后 invalidate。", "我的行程页、行程相关刷新。"],
            ["useTripQuery(id)", "['trip', id]", "GET /api/trips/:id", "enabled=id；staleTime 10s；PATCH 后 setQueryData。", "详情页、分享页。"],
            ["useImportJob(id)", "['import-job', id]", "GET /api/imports/:id", "enabled=id；refetchInterval 1500ms。", "解析页。"],
            ["useDestinations", "['destinations']", "GET /api/destinations", "staleTime 60s。", "目的地页。"],
            ["useExploreRoutes(dest)", "['explore', dest]", "GET /api/explore?dest=...", "enabled=dest；staleTime 60s。", "热门路线页。"],
            ["useFeaturedRoutes", "['featured-routes']", "GET /api/n8n-routes?limit=30", "staleTime 5min。", "首页推荐/灵感/正在看。"],
            ["useTravelInfo(tripId, dayId)", "['travel-info', tripId, dayId]", "POST /api/travel-info", "enabled=tripId+dayId；staleTime 5min。", "Trip DaySection 与地图连线。"],
            ["createImport", "无单独长期 key", "POST /api/imports", "成功后 invalidate ['trips']。", "首页、导入页。"],
            ["toggleFavorite/regenerate/optimizeDay/addSpot/deleteSpot/updateSpot", "['trip', id] + ['trips']", "PATCH /api/trips/:id", "syncTrip 写入单条 Trip，再 invalidate trips。", "详情页和我的行程。"],
            ["deleteTrip", "['trip', id] + ['trips']", "DELETE /api/trips/:id", "removeQueries 单条，invalidate trips。", "我的行程。"],
            ["addExploreRoute", "['trip', id] + ['trips']", "POST /api/explore/add", "直接 Trip 则写缓存；AI job 则返回 jobId。", "Explore、Quiz。"],
        ],
        [1.35, 1.15, 1.35, 1.35, 1.3],
    )

    add_heading(doc, "四、后端 API 工作表", 1)
    add_table(
        doc,
        ["接口", "请求方", "服务端处理", "数据/外部依赖", "返回"],
        [
            ["GET /api/health", "运维/健康检查", "直接返回 ok。", "无。", "{ok:true}"],
            ["GET /api/trips", "my-trips、React Query", "repository.listUserTrips(ownerId)。", "PostgreSQL user_trips 或本地 JSON。", "Trip[]"],
            ["GET /api/trips/:id", "trip/share", "先查用户私有行程，再查 public route。", "Repository。", "Trip 或 404。"],
            ["PATCH /api/trips/:id", "trip/my-trips", "Zod 校验 action；支持收藏、重新生成、优化 Day、增删改 Spot。", "Repository updateTrip + optimizeTrip/optimizeDay。", "更新后的 Trip。"],
            ["DELETE /api/trips/:id", "my-trips", "按 ownerId 删除。", "Repository deleteTrip。", "{ok:true} 或 404。"],
            ["GET /api/routes/public", "公共路线 hook", "解析 destination/tag/budget/q/days 过滤。", "Repository public routes。", "Trip[]。"],
            ["POST /api/routes/:id/save", "公共路线保存", "复制公共路线为用户 private_draft。", "Repository savePublicRoute。", "新 Trip。"],
            ["POST /api/trips/:id/publish", "发布行程", "把 user trip 标为 published_pending，并复制 public route。", "Repository publishTrip。", "更新后的 user Trip。"],
            ["POST /api/imports", "首页/导入页", "创建 import job，异步 processImportJob：DeepSeek 解析、地理编码、优化、保存。", "DeepSeek、Nominatim、Repository、TTLCache。", "ImportJob 初始状态。"],
            ["GET /api/imports/:id", "解析页", "从 TTLCache 读取任务。", "importJobs 内存缓存。", "ImportJob 或 404。"],
            ["POST /api/travel-info", "地图/DaySection", "校验 trip/day，先查内存，再查 route_segments，缺失则 getDayTravelInfo 并保存。", "route_segments、AMap/OSRM/估算路线。", "TravelInfoMap。"],
            ["GET /api/spot-image", "地点详情图", "先查内存，Bing CN，Wikipedia，最后 fallback。", "spotImageCache、Bing、Wikipedia。", "302 重定向图片 URL。"],
            ["GET /api/destinations", "目的地页", "优先从 n8n PostgreSQL 按 destination/country 聚合路线数；失败回退硬编码。", "n8n routes 或 tripPlanner。", "DestinationGroup[]。"],
            ["GET /api/explore", "热门路线页", "读取硬编码路线，再合并 n8n published routes 并映射成 ExploreRoute。", "tripPlanner + n8n routes/days/places 聚合。", "ExploreRoute[]。"],
            ["POST /api/explore/add", "Explore/Quiz", "n8n-* 读取 DB 组装 Trip；profile 命中则同步生成；否则创建 AI import job。", "n8n DB、tripPlanner、DeepSeek、Repository。", "Trip 或 {aiJobId,title}。"],
            ["POST /api/quiz-recommend", "Quiz", "Zod 校验偏好，匹配本地路线，同时异步 generateQuizTrip。", "tripPlanner matchQuizToRoutes、DeepSeek、Repository、TTLCache。", "{dbMatches, aiRoutes}。"],
            ["POST /api/quiz-ai-status", "Quiz 轮询", "按 ids 查 quizAiJobs。", "quizAiJobs 内存缓存。", "{routes}。"],
            ["GET /api/n8n-routes", "首页", "按 status='published' 查询 n8n routes，补 cover_url。", "n8n PostgreSQL。", "FeaturedRoute[]。"],
            ["GET /api/n8n-route-detail", "调试/详情接口", "读取 route + days + places 并转成 Trip。", "n8n routes/itinerary_days/itinerary_places。", "{route, trip}。"],
        ],
        [1.2, 1.0, 1.9, 1.5, 0.9],
    )

    add_heading(doc, "五、核心业务链路", 1)
    add_table(
        doc,
        ["业务链路", "用户动作", "前端", "后端", "数据结果"],
        [
            ["AI 导入攻略", "粘贴链接/文本或上传截图，点开始解析。", "createImport -> /parsing 轮询；截图在首页抽屉会压缩成 base64。", "POST /api/imports 创建 job；parseWithAI 调 DeepSeek；缺坐标用 Nominatim；optimizeTrip；saveUserTrip。", "生成 private_draft Trip，解析页自动进入地图。"],
            ["热门路线加入", "在 /explore 点“查看详情”。", "addExploreRoute(route.id)。", "n8n-*：查 route/days/places 组装 Trip；非 n8n：profile 生成或 AI job。", "Trip 保存到 user_trips；前端进入 /trip 或 /parsing。"],
            ["Quiz 定制推荐", "完成六步偏好问卷。", "POST quiz-recommend；结果页展示 DB 匹配卡和 AI 生成状态。", "matchQuizToRoutes 返回匹配；generateQuizTrip 异步生成并保存。", "用户可添加匹配路线；AI 完成后可查看专属 Trip。"],
            ["地图路线展示", "打开 /trip。", "TripMapView 按 day/spots 初始化 MapLibre；DaySection 拉交通信息。", "/api/travel-info 查缓存或生成 TravelInfo；routePoints 用 polyline 连接点。", "地图点线、Day chips、列表交通时间同步展示。"],
            ["地点详情与导航", "点地图点或地点卡，再点详情/导航。", "SpotModal 用 /api/spot-image 显示图；导航链接调用 amapNavUrl。", "spot-image 查图并 302；导航不走后端。", "弹半页图文详情；外部高德打开步行导航。"],
            ["编辑行程", "详情页点编辑，改时间/标题/描述、加点、删点、优化。", "PATCH action；syncTrip 写入 React Query。", "handleTripPatch 修改 day.spots 并 optimizeDay/optimizeTrip。", "Trip JSON 更新，列表和详情保持同步。"],
            ["收藏/删除/重新生成", "我的行程或详情页点对应按钮。", "toggleFavorite/regenerate/deleteTrip。", "PATCH 或 DELETE；Repository 持久化。", "我的行程刷新；单条缓存更新或移除。"],
            ["分享行程", "详情页点分享。", "/share 读取 Trip，生成卡片和文本摘要。", "不新增服务端接口，仅复用 GET /api/trips/:id。", "Web Share/剪贴板/下载 txt。"],
            ["发布/公共路线复用", "调用 publish 或 save public route。", "useTripActions.publishTrip/savePublicRoute。", "Repository 复制数据，visibility/state 改变。", "个人副本和公共路线分离。"],
        ],
        [1.1, 1.15, 1.6, 1.75, 1.0],
    )

    add_heading(doc, "六、数据源与存储链路", 1)
    add_table(
        doc,
        ["数据类型", "当前来源", "当前存储", "被谁消费", "上线建议"],
        [
            ["用户私人行程", "导入、Quiz AI、Explore 添加、seed。", "PostgreSQL user_trips 或本地 JSON userTrips。", "/my-trips、/trip、/share。", "上线统一 RDS；anonymous_user_id 后续绑定真实账号。"],
            ["公共路线", "seed publicRoutes、publishTrip。", "当前 Repository routes 表或 JSON publicRoutes。", "/api/routes/public、savePublicRoute。", "与 n8n routes 拆名，避免表冲突。"],
            ["n8n 精品路线", "routey-n8n-workflow 每日 DeepSeek 生成。", "n8n PostgreSQL routes、itinerary_days、itinerary_places。", "首页、目的地页、Explore、Explore add。", "加入来源审核和质量校验，不能仅 sourceVerified=true。"],
            ["硬编码路线/目的地", "tripPlanner.ts、extraRoutes.ts、exploreRegionProfiles.ts。", "代码内数组。", "fallback、Quiz 匹配、Explore fallback。", "迁入 explore_routes、destinations、route_profiles。"],
            ["路线段交通", "/api/travel-info 生成。", "内存 travelCache + Repository route_segments。", "地图连线、DaySection 时间距离。", "按 trip/day 或坐标 pair 加过期策略。"],
            ["导入任务状态", "POST /api/imports 创建。", "TTLCache importJobs。", "解析页轮询。", "迁入 import_jobs，服务重启不丢状态。"],
            ["Quiz AI 状态", "POST /api/quiz-recommend 创建。", "TTLCache quizAiJobs。", "Quiz 结果页轮询。", "迁入 quiz_sessions。"],
            ["地点图片", "spot.image 或 /api/spot-image 查图。", "内存 spotImageCache；部分 spot.image 在 Trip JSON。", "SpotModal。", "落 spot_image_cache，并把稳定图片转 OSS。"],
        ],
        [1.1, 1.3, 1.45, 1.2, 1.45],
    )

    add_heading(doc, "七、外部服务与地图链路", 1)
    add_table(
        doc,
        ["服务", "用途", "调用位置", "触发条件", "失败回退"],
        [
            ["DeepSeek", "解析攻略、Quiz 生成、Explore 无 profile 时生成精品路线。", "src/lib/ai.ts；n8n workflow。", "用户导入、Quiz 提交、Explore add fallback。", "导入失败时 createTripFromImport 关键词 fallback；Quiz job 标记 error。"],
            ["Nominatim", "给缺少坐标的 AI spots 补经纬度。", "src/lib/amap-api.ts batchGeocode。", "AI 返回 spot 缺 lat/lng。", "失败保留无坐标 spot；地图不会画该点。"],
            ["MapLibre + CARTO Voyager", "地图底图展示，全球可用且不需要 key。", "src/components/AMapView.tsx。", "TripMapView 初始化。", "map error 后切 SVG SketchMap。"],
            ["高德 URI", "打开外部 APP 导航/搜索。", "amapNavUrl、amapSearchUrl。", "用户点导航按钮。", "无坐标则走关键词搜索。"],
            ["高德 Web API", "中国境内实时 walking/driving/transit 路线。", "src/lib/amap-api.ts。", "ENABLE_LIVE_ROUTE_PLANNING=true 且两点在中国。", "OSRM 或估算路线。"],
            ["OSRM", "海外实时道路 polyline。", "src/lib/amap-api.ts getOsrmRoute。", "ENABLE_LIVE_ROUTE_PLANNING=true 且非中国或高德失败。", "直线估算 polyline。"],
            ["Bing CN Images", "优先找国内可访问的景点图片。", "server.ts fetchBingImage。", "GET /api/spot-image 且缓存未命中。", "Wikipedia。"],
            ["Wikipedia Images", "国际地点图片备用。", "server.ts fetchWikipediaImage。", "Bing 无结果。", "内置 Unsplash fallback。"],
            ["n8n", "自动生成路线资产。", "routey-n8n-workflow.json。", "每日 08:00 schedule。", "前端 fallback 到硬编码路线/目的地。"],
        ],
        [1.0, 1.45, 1.35, 1.45, 1.25],
    )

    add_heading(doc, "八、运行状态与错误处理", 1)
    add_table(
        doc,
        ["模块", "正常状态", "异常/边界", "当前处理", "建议"],
        [
            ["SSR/服务端错误", "TanStack Start 正常返回页面或 API。", "未处理异常、h3 吞掉 SSR error。", "start.ts middleware 和 server.ts normalizeCatastrophicSsrResponse 返回品牌错误页。", "生产接入日志平台和告警。"],
            ["导入任务", "processing -> done/error；progress 0-100。", "DeepSeek key 缺失、模型失败、内容不足。", "AI 失败时 fallback keyword trip；严重失败返回 error。", "任务状态落 DB，记录原始输入和错误。"],
            ["Quiz AI", "generating -> done/error。", "生成超时或失败。", "前端最多轮询约 60 秒；失败显示 error。", "持久化 job，支持重试。"],
            ["地图底图", "MapLibre 加载 CARTO style。", "网络失败、style 加载异常。", "切换 SketchMap SVG fallback。", "国内上线可考虑自托管瓦片/合规底图方案。"],
            ["路线规划", "缓存命中或生成 TravelInfo。", "缺坐标、外部 API 失败。", "缺坐标跳过；失败走 estimateRoute。", "把失败原因写 route_segments metadata。"],
            ["图片", "缓存命中或 302 到图片。", "外链失效、Bing/Wiki 不稳定。", "fallback 到内置图片。", "OSS 转存高频图片。"],
            ["数据源", "DATABASE_URL 存在则 PostgreSQL；否则 JSON。", "RDS 连接失败、schema 不一致。", "当前没有自动降级到 JSON；连接失败会影响 API。", "启动时做 DB healthcheck 和 migration 校验。"],
            ["前端请求", "api<T> 非 2xx 抛 Error。", "接口返回非 JSON、网络失败。", "页面局部 error/toast 或空状态。", "统一 toast/error boundary 体验。"],
        ],
        [1.05, 1.35, 1.35, 1.45, 1.3],
    )

    add_heading(doc, "九、关键文件索引", 1)
    add_table(
        doc,
        ["文件", "职责", "全链路位置"],
        [
            ["src/routes/__root.tsx", "全局 HTML shell、QueryClientProvider、meta/css、404/错误组件。", "所有页面的最外层。"],
            ["src/router.tsx", "创建 TanStack Router 和 QueryClient。", "客户端/服务端路由上下文。"],
            ["src/start.ts", "TanStack Start 实例和 server middleware。", "SSR/请求错误保护。"],
            ["src/server.ts", "所有 /api 分发、AI job、n8n DB、图片、旅行时间、SSR fetch handler。", "后端核心。"],
            ["src/lib/tripStore.ts", "前端 hooks、React Query key、API 封装、写操作缓存同步。", "前端数据总线。"],
            ["src/lib/tripRepository.ts", "Repository 抽象、Memory/File/PostgreSQL 三套实现。", "后端数据访问总线。"],
            ["src/lib/ai.ts / prompts.ts", "DeepSeek 调用、JSON 解析、坐标补齐、Quiz 生成 prompt。", "AI 生成和导入解析。"],
            ["src/lib/amap-api.ts", "地理编码、路线规划、OSRM、估算路线。", "地图路线段生成。"],
            ["src/lib/tripPlanner.ts / extraRoutes.ts", "seed trips、Explore fallback、Quiz 匹配、profile 生成。", "路线资产 fallback。"],
            ["src/components/AMapView.tsx", "MapLibre 地图、SVG fallback、路线线段/点位/高德 URI。", "Trip 地图核心。"],
            ["scripts/serve-aliyun.mjs", "生产 Node HTTP 服务器，静态资源和 server fetch bridge。", "阿里云运行入口。"],
            ["deploy/aliyun/nginx.conf / setup.sh / ecosystem.config.cjs", "Nginx、SSL、PM2 部署。", "上线运维入口。"],
            ["routey-n8n-workflow.json", "自动生成精品路线并写 PostgreSQL。", "路线资产生产入口。"],
            ["database-schema.sql / db/001_route_database.sql", "数据库目标 schema/旧 migration。", "上线前需要统一。"],
        ],
        [1.65, 2.55, 2.3],
    )

    add_heading(doc, "十、上线前全链路验收清单", 1)
    add_table(
        doc,
        ["编号", "验收项", "通过标准"],
        [
            ["1", "首页动态数据", "n8n 有数据时首页灵感/推荐/正在看来自 /api/n8n-routes；DB 断开时 fallback 不白屏。"],
            ["2", "导入攻略", "文本/链接/图片都能创建 job；解析页可轮询；done 后自动进入 /trip?focus=map。"],
            ["3", "热门路线", "目的地页进入 Explore；n8n route 点击后直接生成 user_trip；非 n8n fallback 可进入 AI 解析。"],
            ["4", "地图详情", "所有有坐标的点显示在地图；线段连接点；Day chip 可聚焦；地点详情图可加载；导航能打开高德。"],
            ["5", "编辑与持久化", "收藏、添加、删除、修改 spot、重新生成后刷新页面仍存在。"],
            ["6", "我的行程", "导入/Explore/Quiz 生成的 Trip 都在 /my-trips 出现；删除后不再显示。"],
            ["7", "Quiz", "问卷提交能返回 dbMatches；AI 完成后可跳详情；失败有可理解状态。"],
            ["8", "分享", "分享页能生成摘要，复制/下载可用。"],
            ["9", "数据库", "DATABASE_URL/N8N_DATABASE_URL 指向同一套已迁移 schema；routes 命名冲突已解决。"],
            ["10", "生产部署", "build:aliyun 成功；PM2 启动；Nginx 静态资源缓存和 API 代理正常；/api/health 返回 ok。"],
            ["11", "外部服务", "DeepSeek、地图、图片、n8n key/credential 都配置在服务端或平台凭证，不泄露到前端。"],
            ["12", "来源真实性", "公共精品路线只有存在可审计 source_references 时才标记 sourceVerified。"],
        ],
        [0.6, 2.4, 3.5],
    )

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build_document()
